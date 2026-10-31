"""
Read a task list from CSV, PNG/JPG (screenshot or photo) or PDF.

Returns a pandas DataFrame with the columns the app's table uses plus a list
of plain-language notes. Images and scanned PDFs go through OCR, so the result
is a *draft* that the user should check in the editable table.
"""
from __future__ import annotations

import io
import re
from datetime import date, timedelta
from typing import List, Optional, Tuple

import pandas as pd

COLUMNS = ["subject", "title", "type", "effort_hours", "deadline",
           "start", "fixed", "max_shift_days"]
TYPES = ["assignment", "quiz", "exam", "project", "lab", "tutorial"]
DEFAULT_EFFORT = {"assignment": 6, "project": 15, "tutorial": 3, "lab": 5,
                  "quiz": 3, "exam": 10}

SYNONYMS = {
    "subject": {"subject", "course", "module", "class", "paper", "unit"},
    "title": {"title", "task", "name", "assignment", "work", "description", "topic", "item"},
    "type": {"type", "kind", "category"},
    "effort_hours": {"effort_hours", "effort", "hours", "hrs", "hour", "time", "estimated hours",
                     "est hours", "estimated effort", "effort hours", "est. hours", "workload"},
    "deadline": {"deadline", "due", "due date", "date", "submission", "submission date",
                 "last date", "due on"},
    "start": {"start", "start date", "release", "release date", "start (optional)"},
    "fixed": {"fixed", "fixed (can't move)", "cant move", "can't move"},
    "max_shift_days": {"max_shift_days", "max shift", "max delay", "flexibility"},
}

MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}
_MON = r"(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*"

DATE_PATTERNS = [
    # 2026-10-30 / 2026/10/30
    (re.compile(r"\b(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})\b"), "ymd"),
    # 30-10-2026 / 30/10/26  (day first)
    (re.compile(r"\b(\d{1,2})[-/.](\d{1,2})[-/.](\d{4}|\d{2})\b"), "dmy"),
    # 30 Oct 2026 / 30th October
    (re.compile(r"\b(\d{1,2})(?:st|nd|rd|th)?[\s\-]+" + _MON + r"\.?,?(?:[\s\-]+(\d{4}))?\b", re.I), "d_mon"),
    # Oct 30, 2026 / October 30th
    (re.compile(r"\b" + _MON + r"\.?[\s\-]+(\d{1,2})(?:st|nd|rd|th)?(?:,?[\s\-]+(\d{4}))?\b", re.I), "mon_d"),
]
HOURS_RE = re.compile(r"(\d+(?:\.\d+)?)\s*(?:h|hr|hrs|hour|hours)\b", re.I)
TYPE_RE = re.compile(r"\b(assignment|quiz|exam|project|lab|tutorial|test|viva|report|essay|presentation)s?\b", re.I)
TYPE_MAP = {"test": "quiz", "viva": "exam", "report": "lab", "essay": "assignment",
            "presentation": "project"}


# ---------------------------------------------------------------------------
def _make_date(y: int, m: int, d: int) -> Optional[date]:
    try:
        return date(y, m, d)
    except ValueError:
        return None


def _infer_year(m: int, d: int, today: date) -> Optional[date]:
    cand = _make_date(today.year, m, d)
    if cand is None:
        return None
    if (today - cand).days > 60:       # clearly meant for next year
        cand = _make_date(today.year + 1, m, d)
    return cand


def find_date(text: str, today: date) -> Tuple[Optional[date], Optional[re.Match]]:
    """First date in text and its match object."""
    for pat, kind in DATE_PATTERNS:
        for m in pat.finditer(text):
            g = m.groups()
            if kind == "ymd":
                dt = _make_date(int(g[0]), int(g[1]), int(g[2]))
            elif kind == "dmy":
                y = int(g[2]); y = y + 2000 if y < 100 else y
                dt = _make_date(y, int(g[1]), int(g[0]))
            elif kind == "d_mon":
                mon = MONTHS[g[1][:3].lower()]
                dt = _make_date(int(g[2]), mon, int(g[0])) if g[2] else _infer_year(mon, int(g[0]), today)
            else:
                mon = MONTHS[g[0][:3].lower()]
                dt = _make_date(int(g[2]), mon, int(g[1])) if g[2] else _infer_year(mon, int(g[1]), today)
            if dt:
                return dt, m
    return None, None


def _clean(s: str) -> str:
    return re.sub(r"\s+", " ", s.replace("|", " | ")).strip(" -–—:;,.|")


def parse_text(text: str, today: date) -> Tuple[pd.DataFrame, List[str]]:
    """Turn free text / OCR output into task rows. One task per line that has a date."""
    rows, guessed = [], 0
    for raw in text.splitlines():
        line = raw.strip()
        if len(line) < 6:
            continue
        dl, m = find_date(line, today)
        if dl is None:
            continue
        rest = line[:m.start()] + " | " + line[m.end():]
        hm = HOURS_RE.search(rest)
        effort = float(hm.group(1)) if hm else None
        if hm:
            rest = rest[:hm.start()] + " | " + rest[hm.end():]
        cells = [_clean(c) for c in re.split(r"\s{2,}|\t|\||;", rest)]
        cells = [c for c in cells if c and not re.fullmatch(r"[\W_]+", c)]
        if effort is None:      # a lone number cell such as "8" is the hours
            for c in cells:
                if re.fullmatch(r"\d+(?:\.\d+)?", c) and float(c) <= 100:
                    effort = float(c)
                    cells.remove(c)
                    break
        tm = TYPE_RE.search(line)
        kind = "assignment"
        if tm:
            w = tm.group(1).lower()
            kind = TYPE_MAP.get(w, w)
        if effort is None:
            effort = float(DEFAULT_EFFORT.get(kind, 6))
            guessed += 1
        cells = [c for c in cells if not re.fullmatch(r"\d+", c)]
        if not cells:
            continue
        if len(cells) >= 2:
            subject, title = cells[0], " ".join(cells[1:])
        else:
            subject, title = "General", cells[0]
        rows.append(dict(subject=subject, title=title, type=kind, effort_hours=effort,
                         deadline=dl, start=None, fixed=kind in ("quiz", "exam"),
                         max_shift_days=None))
    notes = []
    if guessed:
        notes.append(f"{guessed} task(s) had no hours listed, so we filled in a typical amount "
                     "for their type. Please adjust them.")
    return pd.DataFrame(rows, columns=COLUMNS), notes


# ---------------------------------------------------------------------------
def _normalise_headers(df: pd.DataFrame) -> pd.DataFrame:
    ren = {}
    for c in df.columns:
        key = str(c).strip().lower()
        for target, names in SYNONYMS.items():
            if key in names and target not in ren.values():
                ren[c] = target
                break
    return df.rename(columns=ren)


def _finish(df: pd.DataFrame, today: date) -> Tuple[pd.DataFrame, List[str]]:
    """Coerce a table with (some of) the standard columns into the app's format."""
    notes: List[str] = []
    df = _normalise_headers(df)
    if "deadline" not in df.columns:
        raise ValueError("Couldn't find a deadline / due-date column.")
    if "title" not in df.columns:
        df["title"] = [f"Task {i + 1}" for i in range(len(df))]
    for c in COLUMNS:
        if c not in df.columns:
            df[c] = None
    df["deadline"] = df["deadline"].map(lambda v: _to_date(v, today))
    df["start"] = df["start"].map(lambda v: _to_date(v, today))
    df["type"] = df["type"].map(lambda v: _norm_type(v))
    df["subject"] = df["subject"].fillna("General").astype(str)
    df["title"] = df["title"].astype(str)
    eff = pd.to_numeric(df["effort_hours"].astype(str).str.extract(r"(\d+(?:\.\d+)?)")[0], errors="coerce")
    missing = eff.isna()
    if missing.any():
        eff[missing] = df.loc[missing, "type"].map(lambda t: float(DEFAULT_EFFORT.get(t, 6)))
        notes.append(f"{int(missing.sum())} task(s) had no hours, so we filled in a typical amount.")
    df["effort_hours"] = eff
    df["fixed"] = df["fixed"].map(lambda v: str(v).strip().lower() in {"1", "true", "yes", "y"}
                                  if v is not None and str(v) != "nan" else False)
    df["max_shift_days"] = pd.to_numeric(df["max_shift_days"], errors="coerce")
    bad = df["deadline"].isna()
    if bad.any():
        notes.append(f"{int(bad.sum())} row(s) had a date we couldn't read and were left out.")
        df = df[~bad]
    return df[COLUMNS].reset_index(drop=True), notes


def _norm_type(v) -> str:
    s = str(v).strip().lower() if v is not None else ""
    s = TYPE_MAP.get(s, s)
    return s if s in TYPES else "assignment"


def _to_date(v, today: date) -> Optional[date]:
    if v is None or (isinstance(v, float) and pd.isna(v)) or str(v).strip() in ("", "nan", "None", "NaT"):
        return None
    if isinstance(v, pd.Timestamp):
        return v.date()
    if isinstance(v, date):
        return v
    dt, _ = find_date(str(v), today)
    return dt


# ---------------------------------------------------------------------------
def _ocr_image(img) -> str:
    import pytesseract
    from PIL import ImageOps
    img = ImageOps.grayscale(img)
    if img.width < 1600:                       # small screenshots OCR better when enlarged
        k = 1600 / img.width
        img = img.resize((int(img.width * k), int(img.height * k)))
    img = ImageOps.autocontrast(img)
    return pytesseract.image_to_string(img, config="--psm 6")


def _read_image(data: bytes, today: date):
    from PIL import Image
    try:
        text = _ocr_image(Image.open(io.BytesIO(data)))
    except Exception as e:  # tesseract missing, corrupt image...
        raise ValueError(f"Couldn't read text from the image ({e}).")
    df, notes = parse_text(text, today)
    return df, ["Read from an image, so some words may be misread - please check the table."] + notes


def _read_pdf(data: bytes, today: date):
    """Read digital or scanned PDFs, with OCR fallback when text parsing fails."""
    import pdfplumber

    notes: List[str] = []
    table_frames = []
    page_text = []
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                # Try common table layouts. Header matching is normalized so
                # variants such as "Due Date" and "Estimated Hours" work.
                try:
                    tables = page.extract_tables(
                        table_settings={"vertical_strategy": "lines", "horizontal_strategy": "lines"}
                    ) or []
                    if not tables:
                        tables = page.extract_tables() or []
                except Exception:
                    tables = page.extract_tables() or []

                for table in tables:
                    if not table or len(table) < 2:
                        continue
                    header = [str(cell or "").replace("\n", " ").strip() for cell in table[0]]
                    frame = pd.DataFrame(table[1:], columns=header)
                    normalized = _normalise_headers(frame)
                    if "deadline" in normalized.columns:
                        table_frames.append(frame)
                page_text.append(page.extract_text(layout=True) or page.extract_text() or "")
    except Exception as exc:
        raise ValueError(f"Couldn't open this PDF ({exc}).")

    # Structured tables are more reliable than free-text parsing.
    if table_frames:
        combined = pd.concat(table_frames, ignore_index=True)
        try:
            df, table_notes = _finish(combined, today)
            if not df.empty:
                return df, table_notes
        except ValueError:
            pass

    text = "\n".join(page_text)
    df, parse_notes = parse_text(text, today) if text.strip() else (pd.DataFrame(columns=COLUMNS), [])
    if not df.empty:
        return df, notes + parse_notes

    # OCR is also attempted for PDFs that contain selectable text but whose
    # layout/parser did not identify any task rows.
    try:
        import pypdfium2 as pdfium
        document = pdfium.PdfDocument(data)
        ocr_pages = []
        for index in range(len(document)):
            page_image = document[index].render(scale=3.0).to_pil()
            ocr_pages.append(_ocr_image(page_image))
        ocr_text = "\n".join(ocr_pages)
        ocr_df, ocr_notes = parse_text(ocr_text, today)
        if not ocr_df.empty:
            notes.append("PDF text extraction found no task rows, so OCR was used. Please verify dates and hours.")
            return ocr_df, notes + ocr_notes
    except Exception as exc:
        if not text.strip():
            raise ValueError(f"Couldn't read this PDF. If it is scanned, ensure Tesseract OCR is installed. Details: {exc}")

    raise ValueError(
        "No tasks were detected in this PDF. Use one task per line or a table with columns "
        "such as Subject, Task, Hours and Due Date. Example: Maths | Assignment 3 | 8 hrs | 30 Oct 2026."
    )


def _read_docx(data: bytes, today: date) -> Tuple[pd.DataFrame, List[str]]:
    """Read Word paragraphs and tables without sending data to an online service."""
    try:
        from docx import Document
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise ValueError(f"Couldn't read this DOCX file ({exc}). Install python-docx.")

    notes: List[str] = []
    frames = []
    for table in doc.tables:
        values = [[cell.text.strip().replace("\n", " ") for cell in row.cells]
                  for row in table.rows]
        if len(values) >= 2 and any(values[0]):
            frame = pd.DataFrame(values[1:], columns=values[0])
            if "deadline" in _normalise_headers(frame).columns:
                frames.append(frame)
    if frames:
        try:
            result, table_notes = _finish(pd.concat(frames, ignore_index=True), today)
            if not result.empty:
                return result, table_notes
        except (ValueError, KeyError):
            pass

    # Paragraphs plus table rows as text: supports ordinary Word task lists.
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip().replace("\n", " ") for cell in row.cells))
    result, text_notes = parse_text("\n".join(lines), today)
    if not result.empty:
        notes.append("Read from a Word document. Please review the extracted task names and dates.")
        return result, notes + text_notes
    raise ValueError("No tasks were detected in this DOCX. Add a due date to each task or use a table with Subject, Task, Hours and Due Date columns.")


def _read_csv(data: bytes, today: date) -> Tuple[pd.DataFrame, List[str]]:
    """Read CSV with common encodings/delimiters; fall back to free-text rows."""
    last_error = None
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            raw = data.decode(encoding)
            frame = pd.read_csv(io.StringIO(raw), dtype=str, sep=None, engine="python",
                                skipinitialspace=True)
            if len(frame.columns) > 1 and "deadline" in _normalise_headers(frame).columns:
                return _finish(frame, today)
            # Some exported CSVs are actually one task per line with no header.
            parsed, notes = parse_text(raw, today)
            if not parsed.empty:
                return parsed, ["CSV had no recognized column headers; interpreted each dated row as a task."] + notes
            return _finish(frame, today)
        except Exception as exc:
            last_error = exc
    raise ValueError(f"Couldn't read the CSV ({last_error}).")


def read_tasks(data: bytes, filename: str, today: date) -> Tuple[pd.DataFrame, List[str]]:
    """Convert supported offline files into the exact DataFrame schema used by the app.

    Supported: CSV, DOCX, PDF, PNG, JPG/JPEG, WEBP and TXT. No external LLM/API is used.
    """
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "csv":
        df, notes = _read_csv(data, today)
    elif ext in ("png", "jpg", "jpeg", "webp"):
        df, notes = _read_image(data, today)
    elif ext == "pdf":
        df, notes = _read_pdf(data, today)
    elif ext == "docx":
        df, notes = _read_docx(data, today)
    elif ext in ("txt", "md"):
        try:
            raw = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            raw = data.decode("cp1252", errors="replace")
        df, notes = parse_text(raw, today)
    else:
        raise ValueError("Supported formats: CSV, DOCX, PDF, PNG, JPG/JPEG, WEBP and TXT.")
    if df.empty:
        raise ValueError("No tasks with readable due dates were found. Add a due date to each task. "
                         "Example: Maths | Assignment 3 | 8 hrs | 30 Oct 2026.")
    # Always return only the app's canonical columns in canonical order.
    return df[COLUMNS].reset_index(drop=True), notes
